# Part of Odoo. See LICENSE file for full copyright and licensing details.

import base64
import io
import zipfile

from openpyxl import load_workbook
from openpyxl.cell.cell import TYPE_ERROR
from openpyxl.utils.exceptions import InvalidFileException

from odoo import fields, models
from odoo.exceptions import UserError
from odoo.fields import Command, Domain
from odoo.tools import escape_psql, remove_accents, split_every

# The Excel file is a data contract defined by the requirements document, not a UI
# text: its headers and status values are kept as literal Spanish constants here.
# Maps the normalized (stripped, lower-cased) Excel header to the (field key on the
# row dict, display label used in error messages).
EXCEL_COLUMNS = {
    'fecha': ('date', "Fecha"),
    'cliente': ('partner_name', "Cliente"),
    'vendedor': ('user_name', "Vendedor"),
    'producto': ('product_name', "Producto"),
    'cantidad': ('quantity', "Cantidad"),
    'valor unitario': ('price_unit', "Valor Unitario"),
    'valor total': ('price_total', "Valor Total"),
    'estado': ('state', "Estado"),
}
STATE_MAPPING = {'borrador': 'draft', 'confirmada': 'confirmed', 'cancelada': 'cancelled'}
MAX_ERRORS_SHOWN = 50
# Maximum number of names per batch search, to keep each SQL query bounded
SEARCH_BATCH_SIZE = 500


class LydSaleRecordImport(models.TransientModel):
    _name = 'lyd.sale.record.import'
    _description = 'Import Sale Records from Excel'

    file = fields.Binary(string="Excel File", required=True, attachment=False)
    file_name = fields.Char(string="File Name")

    def action_import(self):
        self.ensure_one()
        workbook = self._load_workbook()
        sheet = workbook.worksheets[0]
        rows_iterator = sheet.iter_rows()
        try:
            header_row = next(rows_iterator)
        except StopIteration:
            raise UserError(self.env._("The Excel file has no rows.")) from None
        header_index = self._parse_header(header_row)

        rows_data = []
        errors = []
        for row_number, row in enumerate(rows_iterator, start=2):
            if all(cell.value in (None, '') for cell in row):
                continue
            row_data, row_errors = self._parse_row(row_number, row, header_index)
            if row_errors:
                errors.extend(row_errors)
            else:
                rows_data.append(row_data)

        if not rows_data and not errors:
            raise UserError(self.env._("The Excel file has no data rows."))
        if errors:
            raise UserError(self._format_errors(errors))

        # Unique names of the file, {cache key: first spelling found}: the original
        # spelling is kept so that new records are not created in lower case.
        unique_partner_names = {}
        unique_user_names = {}
        unique_product_names = {}
        groups = {}
        for row_data in rows_data:
            for unique_names, field_key in (
                (unique_partner_names, 'partner_name'),
                (unique_user_names, 'user_name'),
                (unique_product_names, 'product_name'),
            ):
                unique_names.setdefault(self._cache_key(row_data[field_key]), row_data[field_key])
            key = (
                row_data['date'],
                self._cache_key(row_data['partner_name']),
                self._cache_key(row_data['user_name']),
                row_data['state'],
            )
            groups.setdefault(key, []).append(row_data)

        Partner = self.env['res.partner']
        Product = self.env['product.product']
        partner_cache, new_partners = self._get_or_create_records(
            Partner,
            unique_partner_names,
            Partner._check_company_domain(self.env.company),
        )
        user_cache, new_users = self._get_or_create_records(
            self.env['res.users'].with_context(active_test=False),
            unique_user_names,
            [('share', '=', False)],
            order='active desc, id',
            prepare_vals_list=self._prepare_user_vals_list,
        )
        product_cache, new_products = self._get_or_create_records(
            Product,
            unique_product_names,
            Product._check_company_domain(self.env.company),
        )

        vals_list = []
        for group_rows in groups.values():
            first_row = group_rows[0]
            line_commands = []
            for row_data in group_rows:
                line_vals = {
                    'product_id': product_cache[self._cache_key(row_data['product_name'])].id,
                    'quantity': row_data['quantity'],
                    'price_total': row_data['price_total'],
                }
                if row_data['price_unit'] is not None:
                    line_vals['price_unit'] = row_data['price_unit']
                line_commands.append(Command.create(line_vals))

            vals_list.append({
                'date': first_row['date'],
                'partner_id': partner_cache[self._cache_key(first_row['partner_name'])].id,
                'user_id': user_cache[self._cache_key(first_row['user_name'])].id,
                'state': first_row['state'],
                'import_file_name': self.file_name,
                'line_ids': line_commands,
            })

        records = self.env['lyd.sale.record'].create(vals_list)
        adjusted_count = len(records.line_ids.filtered('price_unit_adjusted'))

        action = self.env['ir.actions.actions']._for_xml_id('lyd_sale_record.lyd_sale_record_action')
        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': self.env._("Import successful"),
                'message': self.env._(
                    "%(orders)s sales with %(lines)s lines imported. New customers: %(partners)s, "
                    "salespeople: %(users)s, products: %(products)s. "
                    "Unit price adjusted in %(adjusted)s lines.",
                    orders=len(records), lines=len(records.line_ids), partners=new_partners,
                    users=new_users, products=new_products, adjusted=adjusted_count,
                ),
                'type': 'success',
                'next': action,
            },
        }

    def _load_workbook(self):
        """Check the file format (D11) and open it as an Excel workbook."""
        if not (self.file_name or '').lower().endswith('.xlsx'):
            raise UserError(self.env._("Unsupported format. Please upload a file with the .xlsx extension."))
        try:
            content = base64.b64decode(self.file)
            return load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        except (zipfile.BadZipFile, InvalidFileException, KeyError, ValueError) as error:
            raise UserError(self.env._("The file could not be read. It may be corrupted or not a valid .xlsx file.")) from error

    def _parse_header(self, header_row):
        """Normalize the header row and check that all the required columns (D11/EXCEL_COLUMNS) are present."""
        header_index = {}
        for col_index, cell in enumerate(header_row):
            normalized = str(cell.value).strip().lower() if cell.value is not None else ''
            if normalized in EXCEL_COLUMNS:
                header_index[EXCEL_COLUMNS[normalized][0]] = col_index
        missing = [label for field_key, label in EXCEL_COLUMNS.values() if field_key not in header_index]
        if missing:
            raise UserError(self.env._(
                "The Excel file is missing the following columns: %(columns)s.",
                columns=", ".join(missing),
            ))
        return header_index

    def _row_error(self, row_number, column, reason):
        return self.env._('Row %(row)s, column "%(column)s": %(reason)s', row=row_number, column=column, reason=reason)

    def _parse_row(self, row_number, row, header_index):
        """Validate a single data row and return ``(row_data, errors)``.

        ``row_data`` is only meaningful when ``errors`` is empty.
        """
        errors = []
        field_labels = dict(EXCEL_COLUMNS.values())

        def cell_at(field_key):
            return row[header_index[field_key]]

        def add_error(field_key, reason):
            errors.append(self._row_error(row_number, field_labels[field_key], reason))

        for field_key in header_index:
            if cell_at(field_key).data_type == TYPE_ERROR:
                add_error(field_key, self.env._("The cell contains an error value."))

        try:
            date_value = fields.Date.to_date(cell_at('date').value)
        except (ValueError, TypeError):
            date_value = None
        if not date_value:
            add_error('date', self.env._("Invalid date. Expected a date or text in YYYY-MM-DD format."))

        row_data = {'date': date_value}
        for field_key in ('partner_name', 'user_name', 'product_name'):
            row_data[field_key] = str(cell_at(field_key).value or '').strip()
            if not row_data[field_key]:
                add_error(field_key, self.env._("This field is required."))

        for field_key in ('quantity', 'price_total'):
            row_data[field_key] = cell_at(field_key).value
            if not self._is_number(row_data[field_key]) or row_data[field_key] <= 0:
                add_error(field_key, self.env._("Expected a number greater than zero."))

        row_data['price_unit'] = cell_at('price_unit').value
        if row_data['price_unit'] in (None, ''):
            row_data['price_unit'] = None
        elif not self._is_number(row_data['price_unit']):
            add_error('price_unit', self.env._("Expected a number."))

        raw_state = cell_at('state').value
        if raw_state in (None, ''):
            add_error('state', self.env._("This field is required."))
        else:
            row_data['state'] = self._parse_state(raw_state)
            if not row_data['state']:
                add_error('state', self.env._(
                    "Unknown status %(value)r. Expected one of: %(values)s.",
                    value=raw_state, values=", ".join(sorted(STATE_MAPPING)),
                ))

        if errors:
            return None, errors
        return row_data, []

    @staticmethod
    def _is_number(value):
        # bool is a subclass of int: a TRUE/FALSE cell is not a number
        return isinstance(value, (int, float)) and not isinstance(value, bool)

    def _parse_state(self, raw_value):
        value = str(raw_value).strip().lower()
        selection_keys = dict(self.env['lyd.sale.record']._fields['state'].selection)
        return STATE_MAPPING.get(value) or (value if value in selection_keys else False)

    def _format_errors(self, errors):
        shown = errors[:MAX_ERRORS_SHOWN]
        if len(errors) > len(shown):
            shown = shown + [self.env._("...and %(count)s more", count=len(errors) - len(shown))]
        return self.env._(
            "The file could not be imported. %(count)s error(s) were found:\n%(details)s",
            count=len(errors), details="\n".join(shown),
        )

    @staticmethod
    def _cache_key(name):
        """Same normalization as the M1 grouping key: Python's lower() also folds accented
        capitals (Ñ, É...), which PostgreSQL's ILIKE does not do with Odoo's "C" collation."""
        return name.strip().lower()

    def _get_or_create_records(self, model, names, extra_domain=(), order=None, prepare_vals_list=None):
        """Find or create the records of ``model`` named ``names`` with a constant number of queries.

        :param model: empty recordset to search in (with its context, e.g. ``active_test``)
        :param dict names: ``{cache key: name}`` of the records to find
        :param extra_domain: domain added to the name search (company, share...)
        :param str order: search order; the first record found for a name is used
        :param prepare_vals_list: callable returning the create values for a list of
            missing names, ``[{'name': name}, ...]`` by default
        :return: ``({cache key: record}, number of records created)``
        """
        cache = {}
        for batch in split_every(SEARCH_BATCH_SIZE, names.values()):
            # escape_psql: "_" and "%" in a name must match literally, not as SQL wildcards
            domain = Domain.OR(Domain('name', '=ilike', escape_psql(name)) for name in batch) & Domain(extra_domain)
            for record in model.search(domain, order=order):
                cache.setdefault(self._cache_key(record.name), record)

        missing_names = [name for key, name in names.items() if key not in cache]
        if missing_names:
            vals_list = prepare_vals_list(missing_names) if prepare_vals_list else [{'name': name} for name in missing_names]
            for name, record in zip(missing_names, model.sudo().create(vals_list)):
                cache[self._cache_key(name)] = record
        return cache, len(missing_names)

    def _prepare_user_vals_list(self, names):
        return [
            {'name': name, 'login': login}
            for name, login in zip(names, self._generate_user_logins(names))
        ]

    def _generate_user_logins(self, names):
        """Build free, ASCII logins for new salespeople (N2), checking archived users too.

        The logins already taken are read in one query per batch; the ones assigned
        here are reserved too, so two new users of the same file never collide.
        """
        base_logins = ['.'.join(remove_accents(name.strip().lower()).split()) for name in names]
        Users = self.env['res.users'].with_context(active_test=False)
        taken_logins = set()
        for batch in split_every(SEARCH_BATCH_SIZE, set(base_logins)):
            domain = Domain.OR(
                Domain('login', '=', base_login) | Domain('login', '=like', escape_psql(base_login) + '.%')
                for base_login in batch
            )
            taken_logins.update(Users.search_fetch(domain, ['login']).mapped('login'))
        logins = []
        for base_login in base_logins:
            login = base_login
            suffix = 2
            while login in taken_logins:
                login = f"{base_login}.{suffix}"
                suffix += 1
            taken_logins.add(login)
            logins.append(login)
        return logins